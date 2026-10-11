"""Build pitch/SharkSight-pitch.pptx, an editable PowerPoint copy of pitch/index.html.

Run: uv run --no-project --with python-pptx python pitch/build_pptx.py
Speaker notes come from pitch/script.md. Keep the text in sync with index.html.
"""

import re
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
FONT = "Helvetica Neue"

VOID = RGBColor(0x00, 0x00, 0x00)
INK = RGBColor(0xF2, 0xF4, 0xF7)
INK2 = RGBColor(0xB4, 0xBA, 0xC4)
INK3 = RGBColor(0x86, 0x8D, 0x99)
HAIR = RGBColor(0x2E, 0x2E, 0x2E)
SWEEP = RGBColor(0x4D, 0x94, 0xFF)
SWEEP_WASH = RGBColor(0x0A, 0x1E, 0x3A)
LOST = RGBColor(0xF5, 0xA5, 0x24)

W, H = Emu(12192000), Emu(6858000)  # 16:9
IN = 914400
LEFT = int(0.9 * IN)
CW = W - 2 * LEFT

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size, color=INK, bold=False, light=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph a str or list of (text, overrides) tuples."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT + (" Light" if o.get("light", light) else "")
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return box


def chrome(slide, n, total):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = VOID
    text(slide, LEFT, H - int(0.5 * IN), CW, int(0.25 * IN),
         "SharkSight · BMW AI for Product Decision Making · EHL – Grand Finale", 10, INK3)
    text(slide, LEFT, H - int(0.5 * IN), CW, int(0.25 * IN), f"{n} / {total}", 10, INK3, align=PP_ALIGN.RIGHT)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Emu(25400), int(W * n / total), Emu(25400))
    bar.fill.solid()
    bar.fill.fore_color.rgb = SWEEP
    bar.line.fill.background()


def kicker(slide, y, label):
    text(slide, LEFT, y, CW, int(0.3 * IN), label.upper(), 12, SWEEP, bold=True)


def heading(slide, y, t, h=int(1.2 * IN), size=34):
    text(slide, LEFT, y, CW, h, t, size, INK)


def bullets(slide, y, items, w=int(CW * 0.9), size=18):
    paras = [[(b, {"bold": True, "color": INK}), (" " + r, {"color": INK2})] for b, r in items]
    box = text(slide, LEFT, y, w, int(0.5 * IN) * len(items), paras, size, spacing=10)
    for p in box.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(0.32 * IN)))
        pPr.set("indent", str(-int(0.32 * IN)))
        clr = pPr.makeelement(qn("a:buClr"), {})
        srgb = clr.makeelement(qn("a:srgbClr"), {"val": "4D94FF"})
        clr.append(srgb)
        font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
        char = pPr.makeelement(qn("a:buChar"), {"char": "○"})
        for el in (clr, font, char):
            pPr.append(el)
    return box


def flow(slide, y, steps, current):
    x = LEFT
    for i, s in enumerate(steps):
        w = int((0.22 + 0.105 * len(s)) * IN)
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, int(0.45 * IN))
        pill.adjustments[0] = 0.5
        pill.fill.background()
        pill.line.color.rgb = SWEEP if s == current else HAIR
        pill.line.width = Pt(1)
        tf = pill.text_frame
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = s
        r.font.name, r.font.size = FONT, Pt(15)
        r.font.color.rgb = INK if s == current else INK2
        x += w
        if i < len(steps) - 1:
            text(slide, x, y, int(0.4 * IN), int(0.45 * IN), "→", 15, INK3, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x += int(0.4 * IN)


def rule(slide, y, w=CW):
    ln = slide.shapes.add_connector(1, LEFT, y, LEFT + w, y)
    ln.line.color.rgb = HAIR
    ln.line.width = Pt(1)


def columns(slide, y, cols, big=34, small=13, w=CW, pick=None):
    rule(slide, y, w)
    cw = w // len(cols)
    for i, (b, s) in enumerate(cols):
        text(slide, LEFT + i * cw, y + int(0.25 * IN), cw - int(0.2 * IN), int(1.4 * IN),
             [[(b, {"size": big, "light": big > 24, "color": SWEEP if b == pick else INK})], [(s, {"color": INK3})]],
             small, spacing=6)


def logo(slide, x, y, size):
    slide.shapes.add_picture(str(HERE / "assets/logo.png"), x, y, size, size)


def round_photo(path, px=600, zoom=1.08):
    im = Image.open(path).convert("RGB")
    s = min(im.size)
    c = int(s / zoom)
    l, t = (im.width - c) // 2, (im.height - c) // 2
    im = im.crop((l, t, l + c, t + c)).resize((px, px), Image.LANCZOS)
    mask = Image.new("L", (px, px), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, px - 1, px - 1), fill=255)
    out = Image.new("RGBA", (px, px))
    out.paste(im, (0, 0), mask)
    buf = BytesIO()
    out.save(buf, "PNG")
    buf.seek(0)
    return buf


def notes_from_script():
    src = (HERE / "script.md").read_text()
    blocks = re.findall(r"\*\*(\d+) · [^*]+\*\*\n(.+?)(?=\n\n|\Z)", src, re.S)
    return {int(n): body.strip() for n, body in blocks}


SLIDES = []


def slide(fn):
    SLIDES.append(fn)
    return fn


@slide
def title(s):
    logo(s, LEFT - int(0.15 * IN), int(1.35 * IN), int(1.3 * IN))
    text(s, LEFT + int(1.3 * IN), int(1.45 * IN), CW, int(1.2 * IN), "SharkSight", 80, light=True)
    text(s, LEFT, int(2.95 * IN), CW, int(1.6 * IN), ["Solid evidence.", "Sound decisions."], 40)
    text(s, LEFT, int(4.75 * IN), int(CW * 0.7), int(1 * IN),
         "A broad knowledge base built from thousands of raw data points. The AI structures it. "
         "The PM decides what the next BMW must deliver.", 20, INK2)


@slide
def team(s):
    kicker(s, int(0.9 * IN), "The team")
    heading(s, int(1.25 * IN), "Zimbabwe", size=40)
    d = int(2.3 * IN)
    cw = int(3.6 * IN)
    x0 = (W - 3 * cw) // 2
    y = int(2.85 * IN)
    people = [("Musa Berkay Kocabaşoğlu", False), ("Atahan Kılıç", True), ("Utku Kalkanlı", False)]
    for i, (name, me) in enumerate(people):
        cx = x0 + i * cw + cw // 2
        if me:
            halo = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2 - Emu(76200), y - Emu(76200), d + Emu(152400), d + Emu(152400))
            halo.fill.solid()
            halo.fill.fore_color.rgb = SWEEP_WASH
            halo.line.fill.background()
            s.shapes.add_picture(round_photo(HERE / "assets/atahan.png"), cx - d // 2, y, d, d)
            ring = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, y, d, d)
            ring.fill.background()
            ring.line.color.rgb = SWEEP
            ring.line.width = Pt(2)
        else:
            text(s, cx - cw // 2, y - int(0.45 * IN), cw, int(0.35 * IN), "✕ Missing", 16, LOST, bold=True, align=PP_ALIGN.CENTER)
            o = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, y, d, d)
            o.fill.background()
            o.line.color.rgb = LOST
            o.line.width = Pt(2)
            o.line.dash_style = 4  # dash
            text(s, cx - d // 2, y, d, d, "?", 60, INK3, light=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, cx - cw // 2, y + d + int(0.2 * IN), cw, int(0.4 * IN), name, 20,
             INK if me else INK3, bold=me, align=PP_ALIGN.CENTER)


@slide
def problem(s):
    kicker(s, int(1.0 * IN), "The problem")
    heading(s, int(1.4 * IN), "Thousands of customer voices. Many markets. One question: what must the successor do?")
    bullets(s, int(2.85 * IN), [
        ("Evidence is scattered:", "verbatims, customer studies, sales and option lists, plus a fast-moving market."),
        ("Markets are different customers:", "the US, Europe and China must not be averaged away."),
        ("Decisions live for years:", "new evidence keeps arriving, and every decision has to stay traceable."),
    ])
    flow(s, int(5.3 * IN), ["Evidence", "Structure", "Definition", "PM decision"], "Evidence")


@slide
def evidence(s):
    kicker(s, int(1.0 * IN), "Evidence")
    heading(s, int(1.4 * IN), "Every source in one space, always live.", h=int(0.8 * IN))
    columns(s, int(2.55 * IN), [
        ("Voice", "what customers say,\nacross car lines"),
        ("Data", "BMW studies, options\nand sales"),
        ("World", "cited web research with\nsource, date and trust"),
        ("Live", "new data keeps flowing in;\nevery update builds on the last"),
    ])
    bullets(s, int(4.65 * IN), [
        ("Honest data:", "internal figures are checked against public ones, and conflicts are shown, never blended silently."),
    ])


@slide
def structure(s):
    kicker(s, int(1.0 * IN), "Structure")
    heading(s, int(1.4 * IN), "The AI reads everything and builds an evidence graph.", h=int(0.8 * IN))
    flow(s, int(2.45 * IN), ["Classify", "Cluster", "Contextualize", "Requirement areas"], "Requirement areas")
    bullets(s, int(3.4 * IN), [
        ("One shared language", "for comments, studies and research, across the whole portfolio."),
        ("Needs are linked to the world:", "what supports them, what contradicts them, and where rivals compete."),
        ("Change is tested statistically,", "so new data can't fake a trend. Defects go to Quality, not into requirements."),
    ])


@slide
def definition(s):
    kicker(s, int(1.0 * IN), "Definition")
    heading(s, int(1.4 * IN), "From need to a measurable requirement, placed on the car.", h=int(0.8 * IN))
    columns(s, int(2.85 * IN), [
        ("Baseline", "today's BMW"), ("Threshold", "rival parity"),
        ("Target", "projected to launch"), ("Stretch", "delighter"),
    ], big=20, small=14, w=int(CW * 0.75), pick="Target")
    bullets(s, int(4.1 * IN), [
        ("One metric the customer feels,", "with a target ladder per car and market. The AI recommends a rung; the PM picks."),
        ("Every number has a basis:", "cited, projected, assumption or gap. The AI never invents a number to fill a gap."),
        ("Requirements are linked,", "like the roof of a House of Quality. Example: range vs. agility, via mass."),
    ])


@slide
def trust(s):
    kicker(s, int(1.0 * IN), "Built for trust")
    heading(s, int(1.4 * IN), "The AI proposes. The PM decides. Everything is on record.")
    bullets(s, int(3.05 * IN), [
        ("Every requirement traces back", "to the evidence behind it."),
        ("Every decision is audited,", "append-only, with who decided and why."),
        ("Nothing is overwritten:", "each run is a new version, linked to the one before."),
        ("Every model call is logged", "by step and cost."),
    ])


@slide
def vision(s):
    kicker(s, int(1.3 * IN), "The vision")
    heading(s, int(1.7 * IN), "A product brain that keeps learning, across every car line and market.")
    text(s, LEFT, int(3.3 * IN), int(CW * 0.7), int(0.9 * IN),
         "Next: prioritization per market, where linked requirements are bundled and conflicts are decided together.", 20, INK2)
    logo(s, LEFT - int(0.1 * IN), int(4.45 * IN), int(0.75 * IN))
    text(s, LEFT + int(0.75 * IN), int(4.55 * IN), CW, int(0.6 * IN), "SharkSight · Live demo →", 26, light=True)


notes = notes_from_script()
for i, build in enumerate(SLIDES, 1):
    s = prs.slides.add_slide(BLANK)
    chrome(s, i, len(SLIDES))
    build(s)
    if i in notes:
        s.notes_slide.notes_text_frame.text = notes[i]

out = HERE / "SharkSight-pitch.pptx"
prs.save(out)
print(out)
