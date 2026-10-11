"""Study, sales and option-list ingest on synthetic files (no BMW data, no paid calls)."""

import json

import openpyxl
import pytest

from bmw_req.db import connect
from bmw_req.ingest import reference
from bmw_req.ingest.options import OptionItem, OptionPage, build_rows, line_ok, page_section, store_rows
from bmw_req.ingest.sales import ingest_sales
from bmw_req.ingest.study import DIST_LABELS, ingest_study


@pytest.fixture
def conn():
    c = connect(":memory:")
    reference.seed(c)
    return c


def write_study(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CN_EU_2025"
    ws.append([None, "Dataset country"])
    ws.append([None, "China", None, "EU"])
    ws.append(["Overall", "BMW 5 Series G68", "BMW 7 Series G70", "BMW 5 Series G60"])
    ws.append(["Mean", 8.0, 9.0, 7.0])
    ws.append(["Seating", None, None, None])
    ws.append(["Mean", 7.0, 9.5, 8.0])
    ws.append(["Storage", None, None, None])
    ws.append(["Mean", 9.0, 8.5, 9.0])
    us = wb.create_sheet("US_2025")
    dists = {  # question → (G60 dist, G70 dist)
        "Overall exterior styling": ([0, 0, 0.1, 0.2, 0.2, 0.2, 0.3], [0.1, 0.1, 0.1, 0.3, 0.2, 0.1, 0.1]),
        "Cupholders": ([0.2, 0.1, 0.1, 0.3, 0.1, 0.1, 0.1], [0, 0, 0, 0.2, 0.2, 0.3, 0.3]),
    }
    first = True
    for q, (a, b) in dists.items():
        us.append([q, None] + (["BMW 5 Series G60", "BMW 7 Series G70 "] if first else []))
        first = False
        us.append(["Sample total", "% of column, sample base", 1, 1])
        for lab, x, y in zip(DIST_LABELS, a, b):
            us.append([lab, "% of column, sample base", x, y])
        us.append([None])
    wb.save(path)


def test_study_ingest_metrics_and_idempotency(conn, tmp_path):
    p = tmp_path / "F70_G60_G68_G70_customer_studies.xlsx"
    write_study(p)
    assert ingest_study(conn, p) == 3 * 3 + 2 * 2
    assert ingest_study(conn, p) == 0
    assert conn.execute("SELECT COUNT(*) FROM study_metric").fetchone()[0] == 13

    g68 = conn.execute(
        """SELECT m.item, m.mean, m.z_within_column, m.rank_within_column, m.n_items_in_column, s.ref, s.wave, s.text
           FROM study_metric m JOIN source_item s ON s.id = m.source_item_id
           WHERE m.vehicle_model_id = 'G68' AND m.market_code = 'CN' ORDER BY m.rank_within_column"""
    ).fetchall()
    assert [r["item"] for r in g68] == ["Seating", "Overall", "Storage"]   # rank 1 = worst
    assert g68[1]["z_within_column"] == pytest.approx(0.0)
    assert g68[0]["z_within_column"] < 0 < g68[2]["z_within_column"]
    assert g68[0]["n_items_in_column"] == 3 and g68[0]["wave"] == 2
    assert g68[0]["ref"].endswith("#CN_EU_2025:R6:C2")
    assert "G68 (CN)" in g68[0]["text"] and "Q-006" in g68[0]["text"]
    # Market header 'EU' covers G60 even though it's the only EU column here.
    assert conn.execute("SELECT market_code FROM study_metric WHERE vehicle_model_id = 'G60'").fetchone()[0] == "EU"

    cup = conn.execute(
        "SELECT * FROM study_metric WHERE item = 'Cupholders' AND vehicle_model_id = 'G60'"
    ).fetchone()
    assert cup["scale"] == "dist_7" and cup["market_code"] == "US"
    assert cup["top2"] == pytest.approx(0.2)
    assert cup["bottom3"] == pytest.approx(0.4)
    assert cup["net"] == pytest.approx(-0.2)
    assert cup["mean"] == pytest.approx(sum((i + 1) * x for i, x in enumerate([0.2, 0.1, 0.1, 0.3, 0.1, 0.1, 0.1])))
    assert cup["rank_within_column"] == 1
    # G70 header has a trailing space in the real file.
    assert conn.execute("SELECT COUNT(*) FROM study_metric WHERE vehicle_model_id = 'G70' AND market_code = 'US'"
                        ).fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM audit_event WHERE actor_id = 'ingest.study'").fetchone()[0] == 1


def write_sales(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, rows in {
        "F70": [("Europe", "EU", 80, 85, 87), ("China", "CN", 0, 0, 0)],
        "G60_G68": [("Europe", "EU", 110, 115, 127), ("China", "CN", 68, 75, 78), ("USA", "US", 75, 78, 80)],
        "G70": [("Rest of World", "RoW", 5, 7, 9)],
    }.items():
        ws = wb.create_sheet(sheet)
        ws.append(["market", "market_code", "volume_2024", "volume_2025", "volume_2030"])
        for r in rows:
            ws.append(list(r))
    wb.save(path)


def test_sales_ingest_waves_models_and_idempotency(conn, tmp_path):
    p = tmp_path / "sales_volumes.xlsx"
    write_sales(p)
    assert ingest_sales(conn, p) == 6 * 2
    assert ingest_sales(conn, p) == 0
    assert conn.execute("SELECT COUNT(*) FROM sales").fetchone()[0] == 6 * 3

    rows = conn.execute(
        """SELECT sa.vehicle_model_id, sa.year, sa.units, sa.is_forecast, s.wave, s.ref FROM sales sa
           JOIN source_item s ON s.id = sa.source_item_id WHERE sa.market_code = 'CN' AND s.ref LIKE '%G60_G68%'
           ORDER BY sa.year"""
    ).fetchall()
    assert [(r["vehicle_model_id"], r["year"], r["units"], r["is_forecast"], r["wave"]) for r in rows] == [
        ("G68", 2024, 68, 0, 1), ("G68", 2025, 75, 0, 2), ("G68", 2030, 78, 1, 2)]
    assert rows[0]["ref"].endswith("#G60_G68:R3:2024") and rows[1]["ref"].endswith("#G60_G68:R3:2025+2030")
    assert conn.execute("SELECT DISTINCT vehicle_model_id FROM sales WHERE market_code = 'US'").fetchone()[0] == "G60"
    meta = json.loads(conn.execute("SELECT meta_json FROM source_item WHERE kind = 'sales'").fetchone()[0])
    assert "Q-015" in meta["notes"]
    assert conn.execute("SELECT COUNT(*) FROM audit_event WHERE actor_id = 'ingest.sales'").fetchone()[0] == 1


PAGE = """SONDERAUSSTATTUNGEN.
                                SA-Nr.   EUR inkl.
AUSSTATTUNGSPAKETE
COMFORT PAKET                   7VB      2.750,–     □      □
– Sitzheizung vorn und hinten
Lenkrad­heizung                248      350,–       □      □
"""


def item(name, line, **kw):
    base = dict(name_de=name, name_en=name, category="Comfort & Space", availability="optional", package_name=None,
                option_code=None, applies_to=[], price_eur=None, source_line=line)
    return OptionItem(**{**base, **kw})


def test_option_line_verification_and_sections():
    assert line_ok("COMFORT PAKET 7VB 2.750,–", PAGE)                      # whitespace collapsed
    assert line_ok("– Sitzheizung vorn und hinten", PAGE)
    assert line_ok("Lenkradheizung", PAGE)                                  # soft hyphen ignored
    assert not line_ok("Sitzheizung hinten", PAGE)
    assert not line_ok("   ", PAGE)
    assert page_section(PAGE) == "options"
    assert page_section("\n  AUSGEWÄHLTE SERIENAUSSTATTUNGEN.\n") == "standard_equipment"
    assert page_section("i5 M60 xDRIVE.\n") == "model_equipment"
    assert page_section("FARBWELT EXTERIEUR.\n") is None
    assert page_section("TECHNISCHE WERTE.\n") is None


def test_option_rows_dedupe_and_store(conn):
    page = OptionPage(items=[
        item("COMFORT PAKET", "COMFORT PAKET 7VB 2.750,–", availability="package", option_code="7VB", price_eur=2750,
             category="Variants & Packages"),
        item("Sitzheizung vorn und hinten", "– Sitzheizung vorn und hinten", availability="package",
             package_name="COMFORT PAKET"),
        item("Lenkradheizung", "Lenkradheizung 248 350,–", option_code="248", price_eur=350),
        item("Lenkradheizung", "Lenkrad heizung", option_code="248"),           # duplicate → dropped
        item("Massagesitze", "Massagesitze vorn"),                            # invented line → unverified
    ])
    rows = build_rows("G60", "G60_OptionList.PDF", {17: PAGE}, {17: (page, 42)})
    assert len(rows) == 4
    assert [r["meta"]["line_verified"] for r in rows] == [True, True, True, False]
    assert rows[0]["ref"] == "G60_OptionList.PDF#p17:comfort-paket"
    assert rows[1]["text"] == "Sitzheizung vorn und hinten (package, DE price list Jul 2026)"

    assert store_rows(conn, "G60_OptionList.PDF", rows, "fake-model") == 4
    assert store_rows(conn, "G60_OptionList.PDF", rows, "fake-model") == 0
    r = conn.execute("SELECT * FROM source_item WHERE kind = 'option_list' AND external_id = '7VB'").fetchone()
    assert (r["vehicle_model_id"], r["market_code"], r["trust_tier"], r["wave"]) == ("G60", "EU", "T1", 1)
    assert r["quote"] == "COMFORT PAKET 7VB 2.750,–"
    meta = json.loads(r["meta_json"])
    assert meta["price_eur"] == 2750 and "assumption" in meta["market_note"] and meta["llm_usage_id"] == 42
    audit = json.loads(conn.execute("SELECT new_value FROM audit_event WHERE actor_id = 'ingest.options'").fetchone()[0])
    assert audit == {"file": "G60_OptionList.PDF", "items": 4, "line_verified": 3, "llm_model": "fake-model",
                     "llm_usage_ids": [42]}
